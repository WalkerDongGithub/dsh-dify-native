"""AST-load only pure extractors from the fixed official source, no Dify imports."""
import ast, csv, io, json, logging, sys
from pathlib import Path
import chardet, docx, pandas as pd, pypdfium2, webvtt, yaml
from docx.document import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
TextExtractionError = ValueError
logger = logging.getLogger(__name__)
UnsupportedFileTypeError = ValueError
NAMES = ['_extract_text_by_file_extension','_extract_text_by_mime_type','_extract_text_from_plain_text','_extract_text_from_json','_extract_text_from_yaml','_extract_text_from_pdf','parser_docx_part','_extract_text_from_docx','_extract_text_from_csv','_extract_text_from_excel','_extract_text_from_vtt','_extract_text_from_properties']
source = Path(__file__).resolve().parents[1] / 'fixtures/dify-cd47a47-document-extractor.py'
module = ast.parse(source.read_text())
exec(compile(ast.Module(body=[n for n in module.body if isinstance(n, ast.FunctionDef) and n.name in NAMES],type_ignores=[]),str(source),'exec'))
if __name__ == '__main__':
    root=Path(sys.argv[1])
    from openpyxl import Workbook
    d=docx.Document(); d.add_paragraph('Actual Office'); d.add_paragraph('Before 段落');t=d.add_table(rows=2,cols=2)
    t.cell(0,0).text='Header';t.cell(0,1).text='Long Heading';t.cell(1,0).text='line1\nline2';t.cell(1,1).text='42';d.add_paragraph('After');d.save(root/'real.docx')
    w=Workbook();s=w.active;s.append([' Name\nHeader ','Value']);s.append(['Ada',3]);s.append([None,None]);s.append([' spaced ',2.5]);s=w.create_sheet('other');s.append(['Col','N']);s.append(['Hello\nworld',True]);s.append(['formula','=1+2']);w.save(root/'real.xlsx')
    (root/'real.pdf').write_bytes((source.parent/'extraction.pdf').read_bytes())
    cases={'real.pdf':'pdf','real.docx':'docx','real.xlsx':'excel','data.csv':'csv','data.json':'json','data.yaml':'yaml','data.vtt':'vtt','data.properties':'properties'}
    (root/'data.csv').write_text('name,long heading\nAda,"hello\nworld"\n')
    (root/'data.json').write_text('{"z":1,"name":"中文"}')
    (root/'data.yaml').write_text('a: 1\nlabel: 中文\n---\nb: [2, 3]\n')
    (root/'data.vtt').write_text('WEBVTT\n\n00:00:00.000 --> 00:00:01.000\n<v Ada>Hello\n\n00:00:01.000 --> 00:00:02.000\n<v Ada>world\n')
    (root/'data.properties').write_text('# comment\na = b\nx:y\n! keep\nlonely\n')
    expected={name:_extract_text_by_file_extension(file_content=(root/name).read_bytes(),file_extension=Path(name).suffix) for name in cases}
    (root/'expected.json').write_text(json.dumps(expected,ensure_ascii=False))

    route_cases=[]
    route_fn=next(n for n in module.body if isinstance(n,ast.FunctionDef) and n.name=='_extract_text_by_file_extension')
    plain_case=next(c for n in ast.walk(route_fn) if isinstance(n,ast.Match) for c in n.cases if isinstance(c.pattern,ast.MatchOr))
    extensions=[p.value.value for p in plain_case.pattern.patterns]
    for i,ext in enumerate(extensions):
        name=f'plain-{i}{ext}';(root/name).write_text('Plain source 中文\nsecond line\n')
        route_cases.append({'file':name,'expected':_extract_text_by_file_extension(file_content=(root/name).read_bytes(),file_extension=ext)})
    for name,mime in [('data.vtt','text/vtt'),('data.properties','text/properties'),('data.yaml','text/yaml'),('real.pdf','application/pdf'),('real.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document'),('real.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),('data.csv','text/csv'),('data.json','application/json')]:
        target='mime-'+name.replace('.','-')
        (root/target).write_bytes((root/name).read_bytes())
        route_cases.append({'file':{'path':target,'mime_type':mime},'expected':_extract_text_by_mime_type(file_content=(root/target).read_bytes(),mime_type=mime)})
    (root/'route-cases.json').write_text(json.dumps(route_cases,ensure_ascii=False))
